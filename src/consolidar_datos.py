"""
Consolida todos los archivos Excel mensuales de la CMF en dos DataFrames:
- cuadro1: índice de provisiones por banco y tipo de crédito
- cuadro2: calidad de cartera (normal, subestándar, en incumplimiento)

La fila "Sistema Bancario" (índice oficial del sistema publicado por la CMF)
se guarda en un CSV aparte para no mezclarla con los bancos.
"""

import re
from pathlib import Path

import openpyxl
import pandas as pd
import xlrd

RAW_DIR = Path(__file__).parent.parent / "data" / "raw"
OUT_DIR = Path(__file__).parent.parent / "data" / "processed"
OUT_DIR.mkdir(parents=True, exist_ok=True)

NOMBRE_SISTEMA = "Sistema Bancario"

# Filas que no son instituciones reales (totales o separadores)
FILAS_IGNORAR = {"total"}

# Los datos de ambos cuadros arrancan en la fila 15 de Excel (base 0: 14)
FILA_INICIO_DATOS = 14

# Posición (base 0) de cada índice. Es la misma en .xls y .xlsx.
# "encabezados" son (fila, columna, texto esperado) para verificar que la CMF
# no cambió el formato antes de confiar en las posiciones.
CUADRO1 = {
    "hoja": 1,
    "columnas": {
        "indice_total": 3,
        "indice_comercial": 5,
        "indice_consumo": 7,
        "indice_vivienda": 8,
    },
    "encabezados": [
        (8, 3, "provisiones"),
        (11, 5, "comercial"),
        (11, 7, "consumo"),
        (11, 8, "vivienda"),
    ],
}

CUADRO2 = {
    "hoja": 2,
    "columnas": {
        "indice_ev_individual": 3,
        "indice_ev_grupal": 5,
        "indice_cartera_normal": 8,
        "indice_cartera_subestandar": 10,
        "indice_cartera_incumplimiento": 12,
    },
    "encabezados": [
        (9, 3, "individual"),
        (9, 5, "grupal"),
        (9, 8, "normal"),
        (9, 10, "subest"),
        (9, 12, "incumplimiento"),
    ],
}


def fecha_desde_nombre(path: Path) -> pd.Timestamp | None:
    """Extrae año y mes desde el nombre del archivo (ej: cmf_2024_05.xlsx)."""
    m = re.search(r"cmf_(\d{4})_(\d{2})", path.stem)
    if not m:
        return None
    return pd.Timestamp(year=int(m.group(1)), month=int(m.group(2)), day=1)


def limpiar_nombre_banco(nombre: str) -> str:
    """Saca asteriscos y notas al pie de los nombres de bancos.
    Ej: 'Banco de Chile (**)' -> 'Banco de Chile'
    """
    return re.sub(r"\s*\([\*\d]+\)", "", nombre).strip()


def es_banco(valor) -> bool:
    """Filtra filas que no son bancos: totales, nulos o separadores."""
    if not isinstance(valor, str) or not valor.strip():
        return False
    limpio = limpiar_nombre_banco(valor).lower()
    return limpio not in FILAS_IGNORAR and not limpio.startswith("nota")


def leer_hojas(path: Path, indices: list[int]) -> dict[int, list]:
    """
    Devuelve {índice de hoja: filas}, con cada fila como lista de valores.
    Los .xls antiguos se leen con xlrd y los .xlsx con openpyxl. Las hojas se
    buscan por índice porque sus nombres tienen problemas de encoding.
    """
    if path.suffix.lower() == ".xls":
        wb = xlrd.open_workbook(path)
        hojas = {}
        for i in indices:
            if i < wb.nsheets:
                ws = wb.sheet_by_index(i)
                hojas[i] = [ws.row_values(r) for r in range(ws.nrows)]
        return hojas

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        nombres = wb.sheetnames
        return {
            i: [list(r) for r in wb[nombres[i]].iter_rows(values_only=True)]
            for i in indices if i < len(nombres)
        }
    finally:
        wb.close()


def encabezados_validos(filas: list, cuadro: dict) -> bool:
    """Verifica que los encabezados estén donde el cuadro espera."""
    for fila, col, texto in cuadro["encabezados"]:
        celda = filas[fila][col] if len(filas) > fila and len(filas[fila]) > col else None
        if not isinstance(celda, str) or texto not in celda.lower():
            return False
    return True


def extraer_cuadro(filas: list, fecha: pd.Timestamp, cuadro: dict) -> list[dict]:
    """Recorre las filas de bancos y toma los índices según las columnas del cuadro."""
    registros = []

    for row in filas[FILA_INICIO_DATOS:]:
        banco = row[1] if len(row) > 1 else None
        if not es_banco(banco):
            continue

        # Si ya llegamos a las notas al pie, paramos
        if banco.strip().startswith("("):
            break

        registro = {"fecha": fecha, "banco": limpiar_nombre_banco(banco)}
        for nombre, col in cuadro["columnas"].items():
            registro[nombre] = row[col] if len(row) > col else None
        registros.append(registro)

    return registros


def procesar_archivo(path: Path) -> tuple[list, list]:
    """Lee un archivo Excel y devuelve los registros de cuadro 1 y 2."""
    fecha = fecha_desde_nombre(path)
    if not fecha:
        print(f"  Saltando {path.name} — no se pudo leer la fecha del nombre")
        return [], []

    cuadros = (CUADRO1, CUADRO2)
    try:
        hojas = leer_hojas(path, [c["hoja"] for c in cuadros])
    except Exception as e:
        print(f"  Error leyendo {path.name}: {e}")
        return [], []

    resultados = []
    for cuadro in cuadros:
        filas = hojas.get(cuadro["hoja"])
        if filas is None or not encabezados_validos(filas, cuadro):
            print(f"  {path.name}: la hoja {cuadro['hoja']} no tiene el formato esperado, se omite")
            resultados.append([])
        else:
            resultados.append(extraer_cuadro(filas, fecha, cuadro))

    return resultados[0], resultados[1]


def limpiar_df(df: pd.DataFrame) -> pd.DataFrame:
    """Reemplaza '---' por NaN y convierte columnas numéricas."""
    df = df.replace("---", pd.NA)
    for col in df.columns:
        if col not in ("fecha", "banco"):
            df[col] = pd.to_numeric(df[col], errors="coerce")
    return df


def guardar(df: pd.DataFrame, nombre: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Ordena y guarda las filas por banco y las del sistema en CSV separados."""
    df = df.sort_values(["fecha", "banco"]).reset_index(drop=True)
    es_sistema = df["banco"] == NOMBRE_SISTEMA

    df_bancos = df[~es_sistema]
    df_sistema = df[es_sistema]
    df_bancos.to_csv(OUT_DIR / f"{nombre}.csv", index=False)
    df_sistema.to_csv(OUT_DIR / f"{nombre}_sistema.csv", index=False)
    return df_bancos, df_sistema


def main():
    archivos = sorted(RAW_DIR.glob("cmf_*.xl*"))
    print(f"Archivos encontrados: {len(archivos)}\n")

    todos_c1, todos_c2 = [], []

    for i, path in enumerate(archivos, 1):
        print(f"[{i:>3}/{len(archivos)}] {path.name}")
        c1, c2 = procesar_archivo(path)
        todos_c1.extend(c1)
        todos_c2.extend(c2)

    # Armar, limpiar y guardar DataFrames
    df_c1, sis_c1 = guardar(limpiar_df(pd.DataFrame(todos_c1)), "provisiones_por_tipo")
    df_c2, sis_c2 = guardar(limpiar_df(pd.DataFrame(todos_c2)), "calidad_cartera")

    print(f"\nDataset cuadro 1: {df_c1.shape[0]} filas × {df_c1.shape[1]} columnas "
          f"(+ {len(sis_c1)} meses del sistema)")
    print(f"Dataset cuadro 2: {df_c2.shape[0]} filas × {df_c2.shape[1]} columnas "
          f"(+ {len(sis_c2)} meses del sistema)")
    print(f"\nBancos encontrados: {sorted(df_c1['banco'].unique())}")
    print(f"\nRango de fechas: {df_c1['fecha'].min().date()} a {df_c1['fecha'].max().date()}")
    print("\nArchivos guardados en data/processed/")


if __name__ == "__main__":
    main()
